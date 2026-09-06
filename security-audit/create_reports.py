from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

out=Path(__file__).resolve().parent
findings=json.loads((out/'findings.json').read_text())
md='''# KreateKaro Pre Release Security Audit

**Date:** 6 September 2026 (Asia/Calcutta). **Baseline:** commit `366f89f` with a clean tracked working tree. **Assessment:** authorized AI-assisted focused pre-release review. **Target:** https://www.kreatekaro.co/ and the available local repository.

## 1. Executive summary

**RED — DO NOT PUBLICLY RELEASE YET.** Based on these three focused audits, KreateKaro / 3DM is not yet reasonably safe for an initial public release. The original application contained three High-severity findings: weak password storage, executable SVG previews, and unsafe Google account merging. Password storage remains open. The SVG fix passes a local component regression; OAuth fixes pass synthetic tests but require real staging-provider verification. No changes were deployed or committed.

The audit recorded **0 Critical, 3 High, 6 Medium and 3 Low application findings**. Of 12 findings, **3 are FIXED AND VERIFIED locally, 2 are FIXED BUT NEED MANUAL VERIFICATION, and 7 are OPEN**. Nine findings are therefore not fully closed if pending manual verification is included. Original severities are retained after fixes. There are no accepted/deferred or false-positive entries in the counted findings table.

Ownership checks are a positive result: actual handler SQL denied cross-user read/update/rename/delete and did not trust browser-supplied user_id. This does not remove the account security risks. Package-scanner severities are listed separately and are not counted as confirmed deployed vulnerabilities.

## 2. Scope

Reviewed authentication, sessions, Google OAuth, email verification, cloud-design APIs, browser imports/storage, source/configuration, package manifest and pnpm lockfile, available Git history and public response headers. The repository has 84 reachable commits. Pattern scanning inspected 810 current/historical/build text objects; exclusions and limitations are below.

Only two controlled public GET requests were performed directly: `/` and `/api/auth/me`, at approximately 06:08 UTC. An earlier web-fetch attempt failed. Public auth/me returned `200 {"user":null}`. No real account was created, accessed or modified. Google and email APIs were not sent test traffic. No destructive requests, load tests or malicious parser bombs were executed. External documentation/advisory lookups were read-only research.

## 3. Architecture and security context

The implementation is a Vite-built React SPA with TanStack Router, Konva and BabylonJS/Three-related libraries. Cloudflare Pages serves `dist`; Pages Functions under `functions/api` implement authentication and D1-backed storage. Documentation mentions TanStack Start, but the checked build is a client Vite build, not evidence of a separate Node SSR service.

Trust boundaries:

1. Untrusted browser requests and imported files cross into Functions and the client renderer respectively.
2. The session cookie is resolved server-side through D1 before cloud-design access.
3. D1 stores users, sessions, OTP codes and JSON design documents. Browser user IDs are not trusted for ownership.
4. Google token/user-info responses supply identity; matching an email is not proof of control of an existing local account.
5. Resend accepts outbound verification messages. Its response is currently ignored by registration, so delivery failure can strand an account.
6. localStorage and IndexedDB belong to the browser profile, not automatically to a signed-in account.
7. Build/development packages run on developer/CI machines; their advisories are not automatically reachable through Workers.

| Component | Implementation and observed behavior |
|---|---|
| Auth utilities | `functions/_auth-utils.ts`; password digest, UUID sessions, cookies, JSON helper, email sender |
| Register | `POST /api/auth/register`; creates unverified password user and verification OTP |
| Login | `POST /api/auth/login`; password and email_verified checks, new random session |
| Verification | `GET /api/auth/verify-otp?email=&code=`; ten-minute intended lifetime, used flag |
| Current account and logout | `GET /api/auth/me`, `DELETE /api/auth/me`; server lookup and session deletion |
| Google | `GET /api/auth/google`; OAuth initiation/callback, token exchange, user-info, session |
| Cloud design collection | `GET/POST /api/designs`; owned list/create, ten-design count check |
| Cloud design item | `GET/PUT/DELETE /api/designs/[id]`; every query binds authenticated user ID |
| Password reset/resend | No implemented route or frontend recovery flow located; schema/helper comments alone do not establish functionality |
| Custom models | `editor.tsx`, `Sidebar.tsx`, `customAssetDB.ts`, `Preview3D.tsx`; browser files -> IndexedDB/blob URLs -> loader |
| Design import/export | `editor.tsx`; JSON import and local/GLB export; `reportGenerator.ts` uses document library APIs |
| Configuration | `wrangler.toml`, `vite.config.ts`, `public/_headers`, `public/_redirects`, `.gitignore` |

No server upload/conversion endpoint was located. fbx2gltf, obj2gltf and gltf-pipeline are declared packages, not evidence of an internet-accessible converter. Script searches found no request-derived child-process execution. `sharp` and `puppeteer` are development dependencies. Browser asset/texture URLs can cause client-side external requests; no server-side request forgery path was identified in these handlers.

## 4. Methodology and evidence

Read all current Functions handlers and shared utilities, schema, auth client, configuration, storage and relevant import/render paths. Searched sources for SQL, dynamic evaluation, HTML sinks, process creation, environment access, browser storage and external fetches. Reviewed manifest/lockfile and pattern-scanned reachable historical text blobs without emitting secret values.

`security-audit/security.test.ts` imports real Functions and executes their actual SQL through Node SQLite with a small D1 API adapter. It is not a live D1 integration test: Cloudflare transaction/isolate semantics and deployment middleware are not reproduced. Google fetches are mocked. `svg.test.tsx` renders the actual Properties component into JSDOM and uses a harmless local event marker; it does not certify a full browser import flow.

Baseline: 11 checks, 6 expected security failures and 5 passes, retained in `baseline-tests.txt`. Final targeted suite: 15 passes. Full project suite: 20 passes, including the existing five consent tests. Tests documenting OPEN weaknesses intentionally assert observed insecure behavior (identical hashes, unthrottled attempts, accepted malformed schema); a green test runner is not a green security decision.

Production frontend builds completed before and after fixes. Backend TypeScript check passed using `tsc -p functions/tsconfig.json`; `git diff --check` passed (line-ending notices only). Network/Node sandbox failures were retried with approved escalation; the successful scanner result is `dependency-audit-online.json`, not the earlier failed audit file. No live server or browser penetration session was started.

## 5. Audit 1 results

Password hashes, account linking and OAuth state require attention as detailed below. Sessions use crypto.randomUUID-derived IDs, are generated on login, and have HttpOnly/Secure/SameSite=Lax/Path=/ cookies. Logout invalidation passed locally. No client-set session identifier is reused on login. Tokens are stored server-side in plaintext, which means a database read compromise also exposes active sessions; token hashing is future hardening.

OTP email matching and type are bound in SQL. Valid OTP sequential use succeeds once; sequential replay fails. Expiry is corrected locally. Concurrent verification can pass the separate read before both used-flag updates; verification is idempotent today, so this is documented without claiming a duplicate-money/token exploit. A future reset flow must use atomic consumption. Password reset, password-change session revocation and resend are **NOT TESTED / NOT IMPLEMENTED IN THE AVAILABLE ROUTES**.

Design ownership: cross-user GET and PUT return 404; DELETE returns 200 but leaves the other user's row untouched. List contains only the caller's designs. A forged create user_id is ignored. SQL metacharacters in an ID do not broaden access. Current-account identity comes only from the session. No BOLA finding was confirmed.

CSRF: SameSite=Lax and JSON APIs reduce ordinary cross-site cookie-bearing mutation risk. Origin/content-type enforcement is absent; same-site untrusted-subdomain behavior was not tested. Wildcard CORS with credentials is invalid in browsers and does not grant wildcard credentialed access. OAuth login CSRF was a separate confirmed issue, now patched locally. Unsupported methods and duplicate cookies/query parameters were inspected only at source level; deployed Pages fallback semantics are **NOT TESTED**.

## 6. Audit 2 results

All inspected D1 queries use bound values; no user-controlled SQL concatenation was found. The important active-content exception was raw SVG HTML insertion and is locally fixed. Other logo style injections are constant strings. `rootElement.innerHTML` is read, not used as an untrusted write. React text rendering and document-library text generation do not establish an XSS sink by themselves. No current request-derived eval, command execution or filesystem traversal path was found.

Design schemas and browser assets remain insufficiently bounded. Negative dimensions and object-valued elements were accepted in a controlled handler test. Very large/deep input, NaN/Infinity conversion behavior, duplicate JSON keys, unusual filenames and geometry/decompression limits were reviewed for missing controls but not exhaustively exercised. No prototype-pollution exploit was demonstrated; no deep merge gadget was located. Arbitrary JSON storage alone is not proof of prototype pollution.

Custom GLB/GLTF inputs are local browser assets, with a five-asset UI count but no file/geometry budget. No server-side converter, temporary-file or upload path exists in the inspected API. ZIP bombs and format fuzzing were deliberately not executed. Cloud create's non-atomic ten-project cap and list-with-full-JSON response should be fixed alongside request budgets, rather than counted as unrelated findings.

## 7. Audit 3 results

### Secrets

`.env` is ignored and contains keys named GOOGLE_CLIENT_ID, GOOGLE_CLIENT_SECRET and RESEND_API_KEY. The scan identified an actual-format Google secret in ignored `.env:3`; its value is never included. No match for that secret class was found in the scoped reachable-history/build scan. This is not, by itself, evidence of a leaked credential. Rotate if this file was shared, published or exposed elsewhere; routine ignored local storage alone does not establish a rotation incident.

The committed `.env.example` contains only a public analytics identifier slot. Browser source uses VITE_GOOGLE_ANALYTICS_ID; the reviewed Functions obtain secrets from env bindings. Google client IDs, D1 database identifiers and analytics measurement IDs are public identifiers, not authentication secrets. No credential was tested against its provider. The scan is pattern-based, excludes large historical files above 2 MB and binary blobs, and does not cover unreachable/deleted Git objects, external logs or cloud secret stores. Nonmatching arbitrary secret formats may be missed.

### Dependencies

Successful pnpm advisory query: 60 advisory entries (1 critical, 20 high, 29 moderate, 10 low), captured with version paths and links. Practical finding KK-09 groups toolchain debt at Medium; no critical production exploit was validated. [Vitest advisory](https://github.com/advisories/GHSA-5xrq-8626-4rwp) applies to UI/API/browser exposure, not the run-only tests used here. [Vite advisory](https://github.com/advisories/GHSA-fx2h-pf6j-xcff) concerns the Windows development server. See dependency-summary.json for exact affected paths. DOMPurify advisories in the tool tree do not explain the original application XSS, which did not call a sanitizer at that sink.

The available scripts expose a Vite dev server and wrangler deployment command. No public development listener was tested or proven. `latest` dependency ranges and placeholder allowBuilds entries reduce reproducibility; lockfile regeneration/install-script policy needs explicit review. No malicious install script was identified, but all transitive package source/scripts were not audited. Package versions were not silently upgraded during this assessment.

### Deployment and browser data

Public `/` returned 200 with frame/MIME/referrer protections; CSP/HSTS/Permissions-Policy absent in the sampled response. `/api/auth/me` returned user:null, DYNAMIC cache status, wildcard CORS plus credentials and no explicit cache policy. Authenticated public caching and edge/WAF configuration are **NOT TESTED**. No deployment equivalence between current commit and live code is assumed.

Vite source maps are not enabled in configuration; no map publication was established. `.env` is not in public. Broad sensitive-file URL probing, private deployment logs, origin bypass, Cloudflare IAM, DNS/TLS configuration and CI/CD credentials were not tested. SPA fallback means a 200 alone would not prove source-file exposure. Local browser draft/asset data persists across logout (KK-08); session cookies are not stored in localStorage. Analytics loading is consent-controlled in the inspected component; deeper third-party-script behavior was not assessed.

## 8. Findings table

| ID | Severity | Finding | Status |
|---|---|---|---|
'''
for f in findings: md+=f"| {f['id']} | {f['severity']} | {f['title']} | {f['status']} |\n"
md+='\n## 9. Detailed findings\n\nReferences with baseline line numbers refer to commit 366f89f, before remediation; function names identify current code.\n'
for f in findings:
 md+=f"\n### {f['id']} {f['title']}\n\n"
 for label,key in [('Audit category','category'),('Severity','severity'),('Status','status'),('Affected component','file'),('What is wrong','wrong'),('Evidence and observed behavior','evidence'),('Attack scenario','scenario'),('User business and technical impact','impact'),('Recommended fix and work performed','fix'),('Confidence','confidence'),('Automatically validated','auto'),('Retest instructions and expected behavior','retest')]:md+=f"**{label}:** {f[key]}\n\n"
md+='''## 10. Fixes performed

Original findings were saved in ORIGINAL_FINDINGS.md before editing four application files. KK-02 removed HTML insertion; KK-03 stopped implicit account merges and required verified provider email; KK-04 bound OAuth state and removed forwarded callback-origin trust; KK-05 normalized expiration comparison; KK-10 added no-store. Session-cookie parsing was additionally anchored to cookie boundaries as trivial hardening.

No password migration, schema migration, dependency upgrade, cloud configuration change, account modification, secret rotation or deployment occurred. All patches remain local and reviewable in Git. The new account-linking rule intentionally rejects email collisions; an explicit authenticated linking feature can be added later. Secure state-cookie behavior requires HTTPS staging validation. Password KDF work was left open because applying an unvalidated KDF in Workers or dropping legacy hashes could make login unusable.

## 11. Retest results

| Check | Result and limit |
|---|---|
| Baseline security checks | 6 failures / 5 passes, demonstrating original weaknesses |
| Final focused regression | 15 / 15 pass; some tests deliberately document open risks |
| Entire test suite | 20 / 20 pass, three files |
| Backend typecheck | Pass |
| Production frontend build | Pass before and after fixes; large bundle warning is not treated as a vulnerability |
| Whitespace diff check | Pass, line-ending notices only |
| Actual D1 / Google / Resend integration | NOT TESTED |
| Deployed fixes | NOT DEPLOYED; NOT TESTED |

Reproduce: `node node_modules/vitest/vitest.mjs run --config security-audit/vitest.config.ts`, `npm test`, `node node_modules/typescript/bin/tsc -p functions/tsconfig.json`, `npm run build`. Node 22.14 exposes experimental node:sqlite in the test adapter. Use a compatible supported Node environment. Baseline tests are historical evidence and should not be regenerated over their saved log after fixes.

## 12. Known limitations and untested areas

No authenticated public account session, deployed D1 binding or transaction behavior, actual Google authorization-code exchange, Google redirect registration, Resend delivery, password reset, production secret configuration, edge rate-limit rules, WAF, bot protection, or account-history remediation was tested. The audit used local synthetic data and two public read-only responses. JSDOM demonstrates an active event sink but is not a full supported-browser compatibility assessment. No stress, destructive file fuzzing, exhaustive malformed-input matrix or full infrastructure penetration test was performed. No claim of absolute absence of secrets, injection or authorization flaws is made.

## 13. Deferred future security reviews

After initial remediation, perform authenticated two-user staging penetration tests, business-logic/quotas/concurrency review, full import/export and file-format fuzzing in isolation, monitoring and alerting, session/OTP cleanup, incident response, backup/restore, Cloudflare IAM and CI/CD review, privacy/compliance and denial-of-service resilience. Repeat security review after launch and when major features, sharing, billing or server-side converters are added.

## 14. Final public launch recommendation

**RED — DO NOT PUBLICLY RELEASE YET.** The product has a sound ownership pattern, but safe password storage is still a High-severity blocker and authentication changes are not yet proven on staging. The current publicly reachable site has not received these patches.

Release gates:

1. Replace and safely migrate password hashes using a vetted KDF validated on the deployment runtime; review accounts previously auto-linked through Google.
2. Deploy the reviewed SVG/OAuth/expiry/cache changes to controlled staging and pass real Google, D1 and two-account import/save/logout regression tests before production rollout.
3. Establish tested shared login/OTP/email limits and bounded design/asset inputs; update exposed development tooling and resolve the release-relevant advisory paths.

The seven OPEN items should have owners and deadlines; the two pending OAuth items must receive provider-backed verification. Re-evaluate the launch decision after these gates. A later GREEN would mean no known Critical/High blockers in this focused scope, not that the software is completely secure.

## Supporting references

- [OWASP password storage guidance](https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html)
- [Cloudflare Workers Web Crypto](https://developers.cloudflare.com/workers/runtime-apis/web-crypto/)
- [Google OpenID Connect](https://developers.google.com/identity/openid-connect/openid-connect)
- [Google email authority caveats](https://developers.google.com/identity/sign-in/android/backend-auth?hl=en)
- Version-specific advisory URLs and dependency paths are retained in dependency-summary.json.

## Evidence inventory

`ORIGINAL_FINDINGS.md`, `findings.json`, `baseline-tests.txt`, `retest-results.txt`, `all-tests.txt`, `functions-typecheck.txt`, `build-before.txt`, `build-after.txt`, `dependency-audit-online.json`, `dependency-summary.json`, `secret-scan.json`, `public-home-headers.txt`, `public-api-headers.txt`, `public-api-body.txt`, `security.test.ts`, `svg.test.tsx` and `vitest.config.ts`. Secret results contain location/type metadata only. The stakeholder Word document intentionally omits exploit code and terminal logs.
'''
(out/'KreateKaro_PreRelease_Security_Audit.md').write_text(md,encoding='utf8')

doc=Document(); sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.72); sec.left_margin=sec.right_margin=Inches(.8)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',30),('Subtitle',15),('Heading 1',18),('Heading 2',13),('Heading 3',11)]:
 s=doc.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('KreateKaro  |  Security review  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
for r in footer.runs:r.font.size=Pt(9)
doc.core_properties.title='KreateKaro Pre Release Security Audit Report'
doc.core_properties.author='Authorized AI assisted security assessment'
def p(t,style=None):return doc.add_paragraph(t,style)
def h(t):doc.add_heading(t,1)
def sub(t):doc.add_heading(t,2)
def page():doc.add_page_break()
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for i,x in enumerate(headers):t.rows[0].cells[i].text=x
 for row in rows:
  for c,x in zip(t.add_row().cells,row):c.text=x
 for ri,row in enumerate(t.rows):
  for ci,c in enumerate(row.cells):
   c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=c._tc.get_or_add_tcPr(); shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'243746' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));pr.append(shade)
   borders=OxmlElement('w:tcBorders')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders);mar=OxmlElement('w:tcMar')
   for side in ['top','left','bottom','right']:
    e=OxmlElement('w:'+side);e.set(qn('w:w'),'100');e.set(qn('w:type'),'dxa');mar.append(e)
   pr.append(mar)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
    for r in para.runs:r.font.size=Pt(9);r.font.bold=ri==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if ri==0 else '000000')
  trpr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trpr.append(cant)
 head=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(head)
 p('')

p('KreateKaro','Title');p('Pre Release Security Audit Report','Subtitle')
p('Platform  KreateKaro / 3DM')
p('Focused Pre Release Application Security Review')
p('6 September 2026')
p('Prepared using\nAuthorized AI assisted security assessment')
doc.add_paragraph().paragraph_format.space_after=Pt(28)
h('RED')
p('DO NOT PUBLICLY RELEASE YET').runs[0].bold=True
p('The review found important account and browser-content weaknesses. Several fixes are ready locally, but password protection still needs improvement and the Google sign-in changes need confirmation in a controlled release environment.')
p('Ready for public release  NO').runs[0].bold=True
p('This report supports a release decision for founders, management and other stakeholders. The technical evidence and reproducible checks are in the accompanying Markdown report.')
p('No changes were deployed to the public website.')
page()
h('1 Executive summary')
p('KreateKaro is still under development. This review was performed before general customer release to identify the most important security problems while they can still be corrected. It deliberately focused on three areas rather than attempting a complete security certification.')
p('The application checks project ownership on the server. Local tests did not let one customer read, change or delete another customer’s cloud project. However, the review found weak password protection, unsafe Google account matching and a way for a malicious logo to run code in the browser.')
p('Five issues received local fixes. Three passed the relevant local verification; two Google sign-in fixes still need confirmation with the real service. Seven findings remain open. The public site has not received these changes.')
h('2 What was tested')
sub('Audit 1 Account and access security')
p('We checked whether an unauthorized person could sign into or take over an account, misuse a session, or access another customer’s projects. This included passwords, verification codes, Google sign-in and cloud-project ownership.')
sub('Audit 2 Application and data input security')
p('We checked whether manipulated requests, malicious input, uploaded files or other unexpected data could compromise the application or another user. A harmless test demonstrated the unsafe logo preview.')
sub('Audit 3 Platform configuration and software supply chain')
p('We checked for exposed credentials, vulnerable software packages and important production configuration weaknesses. Public checks were limited to the homepage and the signed-out account response.')
h('3 Overall result')
p('RED — Ready for public release: NO. Password storage remains a serious issue. The local fixes also need a controlled deployment and real account testing. Passing the automated tests does not mean all security findings are closed.')
p('Local validation completed: 20 automated checks passed, the production build passed and the backend code check passed. No real customer data was used.')
page()
h('4 Security scorecard')
table(['Area','Status','Risk','Plain English outcome'],[
 ['Account and access','ACTION REQUIRED','High','Ownership checks passed; password storage and release verification remain.'],
 ['Application and input','ACTION REQUIRED','High originally','Logo preview fixed locally; input and asset limits remain.'],
 ['Secrets and dependencies','ACTION REQUIRED','Medium','No committed secret identified in the scoped scan; software updates required.'],
 ['Production configuration','PASS WITH IMPROVEMENTS','Low','Some browser protections exist; headers and deployed fixes need review.']], [1.3,1.25,.85,3.07])
table(['Critical','High','Medium','Low','Locally verified','Pending check','Open'],[['0','3','6','3','3','2','7']],[.7,.55,.7,.5,1.25,1.25,1.52])
p('Severity counts describe the original findings. “Open” excludes the two fixes waiting for manual checks; nine findings are not fully closed. Software-scanner alerts are not added to these application counts.')
h('5 Findings summary')
high_text=[
 ('KK 01 Password protection','HIGH | OPEN','If a database copy is stolen, the stored passwords would be too easy to guess. Replace the storage method and safely migrate existing accounts before release. This change needs testing on the hosting platform.'),
 ('KK 02 Malicious logos and projects','HIGH | FIXED AND VERIFIED LOCALLY','A supplied logo or imported project could run code when previewed, potentially reading or changing the customer’s work. The preview now treats the logo as an image. The local regression passed; deploy and check normal imports.'),
 ('KK 03 Google account matching','HIGH | FIXED BUT WAITING FOR CONFIRMATION','Google sign-in could incorrectly activate or join an existing account. The local fix separates accounts and rejects unsafe matches. Confirm real Google signup and login, and review any accounts previously linked by the old flow.')]
for title,status,text in high_text:sub(title);p(status).runs[0].bold=True;p(text)
page()
h('Findings summary continued')
medium_text=[
 ('KK 04 Google login request checks','FIXED BUT WAITING FOR CONFIRMATION','The old flow did not check which browser began Google login. This could put someone into the wrong account. A local fix binds the request to the browser; real sign-in testing remains.'),
 ('KK 05 Expired sessions and codes','FIXED AND VERIFIED LOCALLY','Expired sessions and email codes could work longer than intended. Corrected date checks now reject them in local tests. Confirm the deployed database behaves the same way.'),
 ('KK 06 Repeated account attempts','OPEN','The application does not limit repeated login, registration or verification attempts. This creates guessing and service-abuse risk. Add shared limits and stronger code generation; existing cloud-level rules were not available for review.'),
 ('KK 07 Project and file limits','OPEN','Unexpected project data is accepted, and uploaded models have no clear size budget. This can break projects or stall the browser. Add validation and practical limits before expanding public use.'),
 ('KK 08 Shared browser privacy','OPEN','Signing out leaves designs and custom assets in the browser. Another person using the same browser profile may see them. Separate account storage and clear private state when people sign out.'),
 ('KK 09 Software package updates','OPEN','The software tree contains known security alerts, mainly involving development and processing tools. Update and retest relevant packages. The scanner’s critical alert does not establish a critical flaw in the public website.')]
for title,status,text in medium_text:sub(title);p('MEDIUM | '+status).runs[0].bold=True;p(text)
page()
h('Findings summary continued')
for title,status,text in [
 ('KK 10 Private response caching','FIXED AND VERIFIED LOCALLY','Private account and project responses did not explicitly prohibit caching. The local change adds that instruction. Confirm it appears after deployment.'),
 ('KK 11 Browser protections','OPEN','The public site has some useful browser protections but lacks others, including a policy that limits executable content. Add and test the missing protections without breaking the editor.'),
 ('KK 12 Account membership privacy','OPEN','Registration tells a requester whether an email already has an account. Use consistent responses and limits to reduce unwanted account discovery.')]:sub(title);p('LOW | '+status).runs[0].bold=True;p(text)
h('6 Fixes completed')
table(['State','Findings'],[
 ['Fixed and verified locally','KK 02 logo preview; KK 05 expiration; KK 10 caching'],
 ['Fixed but waiting for confirmation','KK 03 Google account matching; KK 04 Google request checks'],
 ['Still open','KK 01, KK 06, KK 07, KK 08, KK 09, KK 11, KK 12'],
 ['Accepted or deferred findings','None accepted. Future review topics are listed separately.']], [2.05,4.42])
p('The changes are in the local working copy. They have not been committed or published. Password migration, cloud settings and package updates were not performed. A successful local test cannot confirm real Google, email or production database behavior.')
p('No cross-customer cloud-project access was demonstrated. No real secret appears in this report. An ignored local configuration file contains a Google secret; the scoped history scan did not identify a committed copy. Rotate credentials if they were shared or exposed elsewhere.')
page()
h('7 Release blockers')
p('Are there issues that must be resolved before public release? YES.').runs[0].bold=True
for x in [
 'Replace weak password storage and safely migrate existing accounts on the actual hosting runtime.',
 'Deploy the reviewed fixes to controlled staging and confirm Google sign-in, account collisions, project imports, ownership and expiration. Review accounts affected by the old Google-linking behavior.',
 'Establish tested limits for account attempts, verification emails and project/file inputs. Resolve security alerts in any exposed development or processing tools.']:p(x,'List Bullet')
h('8 What this audit does not guarantee')
p('A security audit reduces risk but cannot prove that software contains zero vulnerabilities. This was a focused first review, not a complete certification. No real customer accounts were tested, and the Google, email and deployed database integrations were not exercised.')
p('Cloud permissions, protective traffic rules, authenticated public workflows, recovery procedures and disruptive file or traffic tests were not tested. A password-reset feature was not found in the available implementation. These gaps are not assumed to be secure.')
h('9 Recommended next security review')
p('After the release blockers are addressed, test two controlled customer accounts in staging and check import/export behavior. Review monitoring, incident response, backups and recovery, cloud permissions, build/deployment security, privacy obligations and resistance to service overload. Test complex file formats in isolation.')
p('Schedule another review after launch and whenever major features are added, especially project sharing, payments or server-side file conversion.')
h('10 Final conclusion')
p('Public launch is not recommended yet. The project has useful ownership controls and several local fixes, but password protection remains a major risk. Complete the release blockers and obtain a fresh release decision based on the deployed, tested version.')
p('A future GREEN decision would mean no known serious blockers remain within the reviewed scope. It would not mean the application is completely secure.')
doc.save(out/'KreateKaro_PreRelease_Security_Audit_Report.docx')
print('Created Markdown report and stakeholder DOCX')
