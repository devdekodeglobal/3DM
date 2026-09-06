from pathlib import Path
from zipfile import ZipFile
from collections import Counter
import json
from docx import Document
p=Path(__file__).resolve().parent
f=json.loads((p/'findings.json').read_text())
d=Document(p/'KreateKaro_PreRelease_Security_Audit_Report.docx')
text='\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
assert len(f)==12
assert Counter(x['status'] for x in f)=={'FIXED AND VERIFIED':9,'FIXED BUT NEEDS MANUAL VERIFICATION':3}
assert 'AMBER' in text
report=p/'KreateKaro_PreRelease_Security_Audit.md'
content=report.read_text(encoding='utf8')
for i in range(1,15):assert f'## {i}.' in content
env_file=p.parent/'.env'
env_lines=env_file.read_text().splitlines() if env_file.exists() else []
for i in range(1,13):assert f'KK {i:02}' in text
for title in ['1 Executive summary','2 What was tested','3 Overall result','4 Security scorecard','5 Findings summary','6 Fixes completed','7 Release blockers','8 What this audit does not guarantee','9 Recommended next security review','10 Final conclusion']:assert title in text,title
with ZipFile(p/'KreateKaro_PreRelease_Security_Audit_Report.docx') as z:
 assert z.testzip() is None
 assert b'PAGE' in z.read('word/footer1.xml')
 for name in z.namelist():
  if name.endswith('.xml'):
   raw=z.read(name).decode()
   for line in env_lines:
    if '=' in line:
     key,value=line.split('=',1);value=value.strip().strip('\"\'')
     if ('SECRET' in key or 'KEY' in key) and len(value)>12:
      assert value not in raw, 'Secret value found in DOCX'
      assert value not in content, 'Secret value found in Markdown'
result={'severity':dict(Counter(x['severity'] for x in f)),'status':dict(Counter(x['status'] for x in f)),'docx_zip':'PASS','all_findings_and_sections':'PASS','page_field':'PASS','secret_value_check':'PASS','visual_layout':'NOT VERIFIED: packaged renderer missing LibreOffice; Word COM rendering stalled and was stopped.'}
(p/'document-checks.json').write_text(json.dumps(result,indent=2));print(json.dumps(result,indent=2))
note='\n\n## Document delivery verification\n\nThe stakeholder DOCX passed structural ZIP/XML, required-section, finding-count, page-number-field and exact local-secret-value exclusion checks. Visual pagination QA could not be completed: the packaged renderer lacked LibreOffice, and a Microsoft Word background render stalled and was stopped. The DOCX should be visually checked in Word before external circulation. No PDF/rendered pages are delivered.\n'
if '## Document delivery verification' not in content:report.write_text(content+note,encoding='utf8')
