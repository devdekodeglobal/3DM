"""Generate the remediation assessment from preserved baseline findings and measured results."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

out=Path(__file__).resolve().parent
baseline=out/'baseline-findings.json'
if not baseline.exists(): baseline.write_bytes((out/'findings.json').read_bytes())
findings=json.loads(baseline.read_text(encoding='utf8'))
changes={
'KK-01': ('Versioned PBKDF2 HMAC SHA256 with 600000 iterations, random 16-byte salts and constant-time comparison. A SHA256 prehash preserves migration compatibility. Successful legacy login upgrades the stored record; a bounded administrative backfill helper also strengthens dormant records without plaintext passwords.', 'Unique salts, valid and wrong passwords, malformed parameters, bounded backfill and legacy login upgrade passed. Actual workerd KDF round-trip passed. Target account CPU allowance and production backfill remain NOT TESTED.', 'functions/_password.ts; functions/_password-migration.ts; functions/api/auth/login.ts'),
'KK-02': ('SVG previews use image-only data URLs instead of raw HTML insertion. Logo sizes and raster dimensions are bounded before decoding. Imported projects are validated before entering editor state.', 'Actual Properties component regression passed: the harmless SVG event marker does not execute. The built editor renders an ordinary draft under enforced CSP. The complete malicious-file browser import path remains a staging check.', 'src/components/editor/Properties.tsx; src/lib/assetValidation.ts; src/routes/editor.tsx'),
'KK-03': ('Google identity lookup uses the provider subject only. Unverified email profiles and collisions with local password accounts are rejected. Password login is disabled for Google-linked accounts, including records linked by the old flow. Verification promotes only the current registration challenge password.', 'Mocked provider tests reject unverified profiles and account collisions. New Google login passes in local workerd with a mocked provider. Real Google lifecycle and historical account/session review remain NOT TESTED.', 'functions/api/auth/google.ts; functions/api/auth/login.ts; functions/api/auth/verify-otp.ts'),
'KK-04': ('OAuth state is bound to a Secure HttpOnly __Host cookie and an expiring one-use D1 challenge. PKCE binds the authorization code to its verifier. Duplicate callback parameters, mismatched cookies, replay and expired challenges are rejected.', 'Missing state, replay and expiration regressions passed. Actual local workerd token exchange checks the PKCE verifier and rejects reused state with a mocked provider. Real HTTPS Google redirect and browser cookie lifecycle remain NOT TESTED.', 'functions/api/auth/google.ts; migrations/0001_security.sql'),
'KK-05': ('Session expiry uses normalized SQL datetime comparison. New registration and OAuth challenges use integer epoch expiry. Logout revokes the server session and verification consumes the challenge transactionally.', 'Expired sessions and codes rejected in regression tests. Actual D1 concurrent verification returns one success and one failure; logout invalidates the session. Real deployment expiry checks remain part of rollout.', 'functions/_auth-utils.ts; functions/api/auth/verify-otp.ts'),
'KK-06': ('Atomic shared D1 limits protect IPs and accounts, with bounded cleanup and Retry-After. Verification codes use rejection-sampled cryptographic randomness and keyed HMAC storage, POST JSON only, ten-minute expiry and one-use consumption. Registration sends through waitUntil and removes the exact challenge after delivery failure.', 'Attempt caps, counter reset, concurrent limits, missing-secret failure, expired/wrong codes and single-use verification passed. Local workerd registration and mocked email delivery passed. Provider delivery, sender configuration and production traffic capacity remain NOT TESTED.', 'functions/_security.ts; functions/_verification.ts; functions/api/_middleware.ts; functions/api/auth/register.ts; migrations/0001_security.sql'),
'KK-07': ('Shared schema and complexity budgets validate API, project import and local draft hydration. Cloud quota is enforced in one atomic insert. Custom model import accepts bounded self-contained GLB only and checks embedded resources, geometry counts and graph structure; logo images have size and pixel limits.', 'Cross-user read/update/delete, forged ownership, malformed JSON, unknown fields, hostile object keys, oversized/nested inputs and unsafe asset fixtures rejected. Real local D1 admits exactly ten of twelve concurrent saves. Ordinary built-editor draft renders. Broad file-format fuzzing and the complete model export matrix remain NOT TESTED.', 'shared/validation.ts; src/lib/assetValidation.ts; functions/api/designs/index.ts; functions/api/designs/[id].ts; src/routes/editor.tsx'),
'KK-08': ('Account-labelled workspace initialization runs before private content mounts. Ownership tokens prevent stale writes; logout/account changes clear localStorage, IndexedDB and mounted editor state. Explicit guest sign-in may preserve the guest draft; unlabelled legacy drafts are cleared.', 'fake-indexeddb regressions verify blob cleanup, account switches and stale-write rejection. Isolated browser smoke test confirms private draft removal and authenticated-editor unmount in two tabs after logout. Important legacy drafts must be exported before rollout.', 'src/lib/workspacePrivacy.ts; src/lib/customAssetDB.ts; src/routes/editor.tsx'),
'KK-09': ('Removed unused conversion, Supabase, Three and developer tooling dependencies. Updated supported direct/transitive packages, pinned direct versions and committed the reproducible pnpm lockfile. Build-script permissions are explicit. A narrow Undici 7.28.0 to 7.29.0 override closes the remaining advisory path.', 'Final pnpm audit reports zero Critical, High, Moderate, Low or informational advisories. All 31 tests, API type check, production build, local workerd and built-editor smoke passed. This is an advisory snapshot, not proof that dependencies contain no vulnerabilities.', 'package.json; pnpm-lock.yaml; pnpm-workspace.yaml; vite.config.ts'),
'KK-10': ('Private API success, error and redirect responses explicitly prohibit caching. Auth client fetches also request no-store. Unnecessary credentialed wildcard CORS was removed.', 'Response regressions and actual local middleware checks confirm no-store and absence of wildcard CORS. Deployment header inspection remains outstanding.', 'functions/_security.ts; functions/_auth-utils.ts; src/lib/authClient.ts'),
'KK-11': ('Enforced static CSP permits the required local editor resources and approved provider hosts without inline JavaScript. Theme initialization moved to an external file. API middleware applies consistent MIME, framing, referrer and permission headers. HSTS excludes preload and includeSubDomains. Unsafe API requests require exact same-origin Origin.', 'Origin forgery and method checks passed in local workerd. The built editor rendered under the configured CSP with zero recorded CSP violations or page errors. Real CDN headers, HTTPS policy and all supported browser/provider/model combinations remain NOT TESTED.', 'public/_headers; public/theme-init.js; functions/_security.ts; functions/api/_middleware.ts; index.html'),
'KK-12': ('Known and unknown email registration requests receive the same neutral 202 acknowledgement. Password hashing occurs before the existence branch and email delivery is asynchronous. Shared limits constrain repeated requests.', 'Synthetic known/unknown response status and body are identical; no real-address enumeration occurred. This reduces obvious enumeration but does not claim mathematically indistinguishable timing.', 'functions/api/auth/register.ts; src/components/editor/AuthModal.tsx'),
}
pending={'KK-01','KK-03','KK-04'}
for f in findings:
 f['baseline_status']=f['status'];f['baseline_fix_note']=f['fix']
 f['status']='FIXED BUT NEEDS MANUAL VERIFICATION' if f['id'] in pending else 'FIXED AND VERIFIED'
 f['fix'],f['retest'],f['current_files']=changes[f['id']]
 f['verification_scope']='Local implementation and synthetic validation only; not deployed'
 f['assessment_date']='2026-09-06'
(out/'findings.json').write_text(json.dumps(findings,indent=2,ensure_ascii=False),encoding='utf8')

md='''# KreateKaro Pre Release Security Audit

**Remediation assessment: 6 September 2026. Branch: `security-audit`. Original baseline: `366f89f`; initial audit commit: `22db909`.** This report covers the current branch changes, local evidence and remaining release checks. It supersedes the original RED recommendation for the local implementation; it does not certify the unchanged public deployment.

## 1. Executive summary

**Overall launch status: AMBER — local remediation passed; public release is conditional.** All twelve original findings have code changes. Nine are fixed and verified within the local test scope; three are fixed but need manual release verification (KK-01, KK-03, KK-04). Zero findings remain without an implementation fix. The three pending findings remain unclosed for release. No finding was accepted as an unresolved risk or silently downgraded.

Original application severities remain **0 Critical, 3 High, 6 Medium and 3 Low**. These are historical finding severities, not twelve currently exploitable vulnerabilities. The separate final dependency scanner reports zero advisories. Automated local gates are GREEN for all three focused audits, with the scope and limitations below. A fully GREEN public-launch decision is not justified yet.

The user selected local validation only. No public account was created or changed, no remote database migration/backfill was run, and these changes were not deployed. Existing deployment, credentials, account history and provider configuration are not inferred from local success.

## 2. Scope

Audit 1 covers password storage, registration, email verification, sessions, Google identity and cloud-project authorization. Audit 2 covers request schemas, hostile content, custom assets, project import and shared-browser privacy. Audit 3 covers dependency/configuration changes, secret handling and HTTP security headers.

The baseline included 84 reachable commits and a pattern scan of 810 current/historical/build text objects. That scan found no confirmed committed secret; an ignored local environment file contained credentials. The full historical scan is baseline evidence, not a newly repeated guarantee. New reports are checked for exact local-secret-value exclusion. Only two baseline public GETs, homepage and signed-out `/api/auth/me`, were sampled. No new production probing was required for this local follow-up.

## 3. Architecture and security context

The app is a Vite-built React client with TanStack Router, Konva and BabylonJS. Cloudflare Pages serves the client and Pages Functions implement D1-backed APIs. Same-origin browser requests cross a server boundary where sessions and ownership are checked; browser-provided user IDs never establish authorization. Untrusted project and model data cross a separate client rendering boundary.

Passwords use versioned salted PBKDF2. Registration holds a pending password hash and keyed verification code in a challenge table; the current challenge is consumed in a transactional user mutation. OAuth uses Google subject identity, browser-bound state and PKCE. API middleware enforces method/origin policies and D1 rate limits. Private browser storage is initialized under an account ownership token before rendering.

| Route | Current protocol |
|---|---|
| `/api/auth/register` | POST JSON; neutral acknowledgement and pending verification challenge |
| `/api/auth/verify-otp` | POST JSON only; keyed code digest, expiry, attempt limits, one use |
| `/api/auth/login` | POST JSON; password verification and session creation |
| `/api/auth/google` | GET initiation/callback; state cookie, PKCE, subject identity |
| `/api/auth/me` | GET current account; DELETE revokes session |
| `/api/designs` | GET owned summaries; POST bounded design with atomic quota |
| `/api/designs/[id]` | GET, PUT, DELETE with server-side ownership checks |

There is no implemented password-reset endpoint and no server upload/conversion service in the reviewed code. Removed conversion packages do not represent a claim that such a production service was tested.

## 4. Methodology and evidence

Real handler code runs in synthetic Node SQLite tests and, separately, in an isolated Miniflare/workerd instance with an in-memory D1 database. Local email and Google responses are mocked; unexpected outbound worker requests are rejected. No local production database or real account fixtures are read. The browser smoke uses the actual production build, configured security headers, a fresh browser profile, synthetic account responses and a loopback server; external browser requests are blocked.

Tests cover regressions and bounded concurrency rather than load attacks. No password cracking, destructive parser bombs, real-address enumeration or provider abuse was performed. Unit mocks and local D1 do not prove production provider integration or CPU capacity.

## 5. Audit 1 results

**Local gate GREEN / PASS. Release assurance AMBER.** Password salts, work factor, wrong-password handling, legacy upgrades and bounded backfill passed. Account/IP limits, verification expiry, pre-registration password replacement, one-use verification and server-side logout passed. Cross-user CRUD and forged ownership attempts were rejected. Mocked Google identity, collision, state and PKCE checks passed, including actual local workerd exchange and state replay.

Real Google sign-in and email delivery are NOT TESTED. Target KDF CPU allowance, dormant password backfill and historical Google-linked account/session review are NOT TESTED. These gaps explain the three pending finding statuses and the release checklist.

## 6. Audit 2 results

**Local gate GREEN / PASS.** The unsafe SVG HTML insertion is removed. Project/API validation enforces object shape, size, depth, numeric and element budgets. Custom GLB checks reject externally linked, compressed, oversized and structurally unsafe fixtures before loading. Logo uploads have byte and raster-pixel budgets. D1 admitted ten of twelve concurrent cloud saves, enforcing the free-plan quota atomically.

Browser storage tests cover private blob cleanup, account changes, stale writes and explicit guest draft preservation. The built editor rendered a valid draft under CSP and removed private draft state on logout across two tabs. Full model-format fuzzing, every import/export variant, mobile browsers and multi-device behavior are NOT TESTED. Old unlabelled drafts are deliberately cleared on first secure initialization; export important drafts before rollout.

## 7. Audit 3 results

**Local gate GREEN / PASS. Deployment assurance AMBER.** Unused dependencies were removed, direct versions pinned, vulnerable packages updated and pnpm build-script policies made explicit. The final advisory scan reports zero across all severities, down from sixty baseline scanner entries. Scanner counts are separate from application findings and do not establish production exploitability.

API headers consistently prohibit caching and remove wildcard credentialed CORS. Static CSP and HSTS are configured, and unsafe methods require exact same-origin requests. The built-editor smoke observed zero CSP violations or page errors. Public CDN/header state, cloud secret settings, approved email sender, cloud permissions and production traffic limits are NOT TESTED in this follow-up.

## 8. Findings table

Statuses below describe local verification. Pending entries need the additional release checks specified in their details.

| ID | Original severity | Audit | Current status | Finding |
|---|---|---|---|---|
'''
for f in findings:md+=f"| {f['id']} | {f['severity']} | {f['category']} | {f['status']} | {f['title']} |\n"
md+='\n## 9. Detailed findings\n\nOriginal evidence is retained explicitly as baseline evidence. Original file references may refer to pre-remediation line numbers. Current files are listed separately.\n'
for f in findings:
 md+=f"\n### {f['id']} {f['title']}\n\n**{f['severity']} | {f['status']} | {f['confidence']}**\n\n"
 for label,key in [('Baseline location','file'),('Baseline weakness','wrong'),('Baseline evidence','evidence'),('Plausible scenario','scenario'),('Impact','impact'),('Current implementation files','current_files'),('Fix performed','fix'),('Retest and remaining verification','retest')]:md+=f"**{label}:** {f[key]}\n\n"
md+='''## 10. Fixes performed

All twelve findings received implementation fixes. Additive `migrations/0001_security.sql` creates rate-limit and pending challenge tables; it does not drop users or designs. `functions/_password-migration.ts` provides a bounded compare-and-update helper for administrative backfill. It is not an exposed HTTP endpoint, and no production backfill was executed.

Deploy frontend and API together because verification changed from GET to POST and cloud lists now return summaries fetched separately for loading. Configure SECURITY_SECRET and an approved RESEND_FROM before release. The password migration and account/session review must target the intended database explicitly. See `RELEASE_STEPS.md` for the ordered procedure and recovery limitations.

Budgets are 1 MiB project JSON, 500 total elements, 100-metre space dimensions, 120-character design names, 10 cloud designs, 25 MiB self-contained uncompressed GLB, 256 KiB logos and bounded embedded image dimensions. These are intentional product restrictions. Compressed/external custom models require conversion before import.

## 11. Retest results

| Validation | Measured result | Evidence |
|---|---|---|
| Full automated suite | PASS, 31 tests across 5 files | `all-tests-final.txt` |
| API TypeScript | PASS, no diagnostics | `typecheck-final.txt` |
| Deployable Pages Functions bundle | PASS, compiled locally with Wrangler | `functions-build-final.txt` |
| Production frontend build | PASS; existing large-chunk warning remains | `build-final.txt` |
| Actual local workerd and D1 | PASS: KDF, middleware, single-use verification, login, atomic quota, logout, mocked OAuth PKCE/replay | `worker-retest.txt` |
| Built-editor browser smoke | PASS: valid draft, enforced CSP, no page errors, two-tab sign-out cleanup | `browser-retest.txt` |
| Final package advisories | PASS, all severity counts zero | `dependency-retest.json` |

The build completed in about six minutes on this host; bundle-size warnings are performance observations, not security findings. KDF host wall time is not billed production CPU. Test-only routes and mock credentials reside under `security-audit`, outside Pages Functions and public assets.

## 12. Known limitations and untested areas

NOT TESTED: real provider login/email delivery; target database migrations and password backfill; historical affected-account/session review; production CPU allowance, rate-limit tuning and traffic capacity; deployed authenticated workflows and CDN headers; cloud IAM/WAF/secrets/CI permissions; complete model import/export compatibility; password reset, backups, disaster recovery, monitoring and incident response. The browser smoke uses synthetic APIs, not real provider cookies. JSDOM SVG tests do not replace a full malicious-file browser workflow.

No production state changed. Local pass results cannot be presented as a GREEN certification of the existing public deployment. Dependency scans are time-specific. The baseline secret pattern scan has coverage limits and cannot rule out credentials shared elsewhere.

## 13. Deferred future security reviews

Deeper business-logic penetration testing, advanced file-format fuzzing, denial-of-service resilience, monitoring and incident response, backup/recovery exercises, cloud/build permissions and privacy obligations need a subsequent review. Password recovery requires a separately reviewed implementation. Repeat assessment after project sharing, payments or server-side conversion are added.

## 14. Final public launch recommendation

**AMBER.** Proceed to a controlled release validation stage; do not describe the public launch as fully cleared yet. The local implementation has no remaining unimplemented finding from this focused audit, but three findings remain pending manual assurance. A GREEN decision requires evidence for these release gates:

1. Apply the additive migration, configure secrets and approved sender, backfill legacy passwords, review previously Google-linked accounts and revoke affected old sessions using the controlled procedure. Verify the KDF against the target CPU allowance.
2. Run real Google and email lifecycle checks on HTTPS staging, including returning/new users, cancellation, account collisions, delivery failures and code expiry.
3. Deploy the reviewed frontend/API together and verify actual headers, two-account cloud operations, shared-browser cleanup and representative file imports/exports on supported browsers.

A security review reduces risk but cannot prove zero vulnerabilities. No current finding was marked safe merely to obtain a desired color.

## Supporting references

- [OWASP password storage guidance](https://cheatsheetseries.owasp.org/cheatsheets/Password_Storage_Cheat_Sheet.html)
- [Cloudflare Workers Node crypto support](https://developers.cloudflare.com/workers/runtime-apis/nodejs/crypto/)
- [Google OpenID Connect](https://developers.google.com/identity/openid-connect/openid-connect)

## Evidence inventory

Current: `findings.json`, `all-tests-final.txt`, `typecheck-final.txt`, `build-final.txt`, `dependency-retest.json`, `worker-retest.txt`, `browser-retest.txt`, security test sources and `RELEASE_STEPS.md`. Historical: `baseline-findings.json`, `ORIGINAL_FINDINGS.md`, `baseline-tests.txt`, `dependency-audit-online.json`, `secret-scan.json` and the two public header samples. Earlier retest logs are historical; final logs above are authoritative.
'''
(out/'KreateKaro_PreRelease_Security_Audit.md').write_text(md,encoding='utf8')

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.72);sec.left_margin=sec.right_margin=Inches(.8)
normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(11)
normal.paragraph_format.space_after=Pt(7);normal.paragraph_format.line_spacing=1.08
for name,size in [('Title',30),('Subtitle',15),('Heading 1',18),('Heading 2',13)]:
 s=doc.styles[name];s.font.name='Calibri';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
 s.paragraph_format.space_before=Pt(10);s.paragraph_format.space_after=Pt(6)
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('KreateKaro  |  Security remediation review  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='KreateKaro Pre Release Security Audit Report'
doc.core_properties.author='Authorized AI assisted security assessment'
def p(t,style=None):return doc.add_paragraph(t,style)
def h(t):doc.add_heading(t,1)
def sub(t):doc.add_heading(t,2)
def page():doc.add_page_break()
def table(headers,rows,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for c,w in zip(t.columns,widths):c.width=Inches(w)
 for c,v in zip(t.rows[0].cells,headers):c.text=v
 for row in rows:
  for c,v in zip(t.add_row().cells,row):c.text=v
 for i,row in enumerate(t.rows):
  row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
  for c,w in zip(row.cells,widths):
   c.width=Inches(w)
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EDF0' if i==0 else 'FFFFFF');c._tc.get_or_add_tcPr().append(shade)
   for para in c.paragraphs:
    para.paragraph_format.space_after=Pt(5)
    for r in para.runs:r.font.size=Pt(10);r.bold=i==0
 t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
 p('')

p('KreateKaro','Title');p('Pre Release Security Audit Report','Subtitle')
p('Remediation and release decision');p('6 September 2026');p('Prepared for founders and product stakeholders')
p('Authorized AI assisted application security assessment')
h('AMBER')
p('Local checks passed   Public release remains conditional').runs[0].bold=True
p('The twelve original findings now have implementation fixes. Nine are verified within the local review scope; three need further release verification. No code finding remains without a fix, and the dependency scanner reports zero advisories.')
p('This assessment covers the security-audit branch. It does not certify the unchanged public deployment. The user selected local validation only; no production account, database or website was changed.')
p('Use this report to decide the next release steps. Reproducible technical evidence is in the accompanying Markdown report.')
page()
h('1 Executive summary')
p('The original review found three serious weaknesses involving password protection, unsafe logos and Google account matching, alongside nine other findings. The branch now includes fixes for all twelve. Tests confirm important controls work locally, including project ownership, verification code limits and shared-browser cleanup.')
p('Public launch remains conditional because the real Google and email services, production password migration and target hosting settings have not been verified. These are release gates, not failed local tests.')
h('2 What was tested')
for title,body in [('Audit 1 Account and access security','Passwords, account registration, verification codes, sessions, Google account matching and ownership of cloud projects.'),('Audit 2 Application and input security','Logo previews, project and model validation, limits on saved designs, private browser storage and sign-out behavior.'),('Audit 3 Configuration and dependencies','Software package advisories, secure response headers, browser content policy, secret handling and reproducible package configuration.')]:sub(title);p(body)
h('3 Overall result')
p('AMBER. All three local audit gates pass. All 31 automated tests, the API code check, the production build and the local Workers/database integration checks pass. A separate isolated browser test confirms the built editor works under its content policy and clears private draft state across two tabs at sign-out.')
p('Ready for unrestricted public release   NOT YET').runs[0].bold=True
page()
h('4 Security scorecard')
table(['Area','Local result','Remaining release assurance'],[
 ['Account and access','PASS','Real providers, password backfill, old linked accounts and hosting CPU allowance.'],
 ['Application and input','PASS','Representative imports/exports and supported-browser release checks.'],
 ['Configuration and dependencies','PASS','Apply secrets/migration and verify actual deployment headers.']], [1.65,1.05,3.97])
table(['Critical','High','Medium','Low'],[['0','3','6','3']],[1.65,1.65,1.65,1.72])
p('Counts retain the original finding severities. Nine findings are fixed and verified locally, three are fixed but waiting for manual confirmation, and zero remain without an implementation fix. The three pending findings are still unclosed for release. The separate dependency scan has zero advisories.')
h('5 Findings summary')
summaries=[
 ('KK 01 Password protection','HIGH | WAITING FOR CONFIRMATION','Passwords now use salted, deliberately expensive hashing. Local runtime and migration tests pass. Confirm the hosting CPU allowance and strengthen all old database records before release.'),
 ('KK 02 Malicious logos','HIGH | VERIFIED LOCALLY','Logo previews now treat SVG content as an image. The regression test confirms the old executable preview behavior is removed. Complete representative upload checks during release validation.'),
 ('KK 03 Google account matching','HIGH | WAITING FOR CONFIRMATION','Google accounts no longer silently merge with password accounts. Unsafe identities and collisions are rejected. Real Google testing and review of historically linked accounts remain required.'),
 ('KK 04 Google login request checks','MEDIUM | WAITING FOR CONFIRMATION','Login requests are bound to the initiating browser and can be used only once. Local provider mocks pass; confirm the actual HTTPS provider and cookie lifecycle.')]
for title,status,body in summaries:sub(title);p(status).runs[0].bold=True;p(body)
page()
h('Findings summary continued')
for title,status,body in [
 ('KK 05 Expiration','MEDIUM | VERIFIED LOCALLY','Expired sessions and codes are rejected. Verification consumes the code transactionally, and sign-out revokes the server session.'),
 ('KK 06 Repeated account attempts','MEDIUM | VERIFIED LOCALLY','Shared request limits, strong code generation, protected code storage and single-use verification reduce guessing and email abuse. Approved sender and real delivery still need release checks.'),
 ('KK 07 Project and file limits','MEDIUM | VERIFIED LOCALLY','The app rejects oversized or malformed project data and unsafe custom model fixtures. Concurrent saves cannot exceed the ten-project limit. Supported custom formats now have explicit restrictions.'),
 ('KK 08 Shared browser privacy','MEDIUM | VERIFIED LOCALLY','Sign-out and account changes clear private browser state. Both the storage regressions and a two-tab built-editor test pass. Export important old drafts before rollout.'),
 ('KK 09 Software packages','MEDIUM | VERIFIED LOCALLY','Unused packages were removed, versions pinned and vulnerable packages updated. The final advisory scan reports zero alerts; tests and the build pass.'),
 ('KK 10 Private response caching','LOW | VERIFIED LOCALLY','Account and project responses explicitly prohibit caching. Actual local middleware tests confirm this behavior.'),
 ('KK 11 Browser protections','LOW | VERIFIED LOCALLY','Content, framing, transport and permission protections are configured. The built editor passes the local content-policy smoke test. Actual deployment headers remain a release check.'),
 ('KK 12 Account membership privacy','LOW | VERIFIED LOCALLY','Registration returns the same acknowledgement for known and unknown emails. Shared limits reduce repeated probing. No real customer addresses were tested.')]:sub(title);p(status).runs[0].bold=True;p(body)
page()
h('6 Fixes completed')
table(['State','Findings'],[
 ['Fixed and verified locally','9: KK 02, KK 05, KK 06, KK 07, KK 08, KK 09, KK 10, KK 11, KK 12'],
 ['Fixed but waiting for confirmation','3: KK 01, KK 03, KK 04'],
 ['Still open without a fix','0'],
 ['Accepted or deferred findings','0; future review topics are separate.']], [2.3,4.37])
p('The fixes are isolated on security-audit. No change was merged into main or deployed by this remediation. Local database tests use synthetic accounts and isolated data. Configuration secrets are excluded from the report.')
h('7 Release blockers')
p('Must anything be resolved before public release? YES.').runs[0].bold=True
for item in [
 'Apply the database migration, configure required secrets and approved email sender, strengthen dormant password records, review old Google-linked accounts and revoke affected sessions. Confirm hosting capacity for password hashing.',
 'Complete real Google and email lifecycle checks on HTTPS staging, including new and returning users, account collisions, cancellation, delivery failures and expired codes.',
 'Deploy the reviewed client and API together, then verify real headers, two-account cloud operations, sign-out cleanup and representative imports/exports on supported browsers.']:p(item,'List Bullet')
page()
h('8 What this audit does not guarantee')
p('A security audit reduces risk but cannot prove that software contains zero vulnerabilities. This was a focused first review, not complete certification. Passing local tests does not certify the existing public deployment.')
p('Real provider behavior, production database changes, infrastructure permissions, traffic capacity, complete model-format compatibility, recovery and incident response were not tested. No password-reset feature is included. Mocked browser accounts do not reproduce provider authentication.')
h('9 Recommended next security review')
p('After the release gates pass, review monitoring, incident response, backups and recovery, cloud and build permissions, privacy requirements, service-overload resilience and deeper file-format/business-logic testing. Repeat the review when project sharing, payments or server-side file conversion are introduced.')
h('10 Final conclusion')
p('The local security posture has improved substantially: all original findings have code fixes and all measured local gates pass. The remaining major uncertainty is whether the fixes and account migrations work correctly with the real deployment and providers.')
p('Proceed with controlled release validation. Approve a GREEN public launch only after the documented release gates have evidence. Until then the recommendation remains AMBER.')
doc.save(out/'KreateKaro_PreRelease_Security_Audit_Report.docx')
print('Updated Markdown, findings and stakeholder DOCX: AMBER; 9 verified, 3 pending, 0 without a fix.')
